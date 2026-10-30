"""Regressions for the three-page Word/two-page PDF overflow mismatch."""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import base64
import sys
import zipfile
import pytest
from docx import Document
from lxml import etree
from pypdf import PdfReader

ROOT=Path(__file__).parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'revision-service'),str(Path(__file__).parent)]
from residential_notice import render_residential_notice, fields_for_case, NoticeInputError
from residential_word_snapshot import saved_word_case, render_saved_word_pdf
from residential_pdf import render_residential_pdf
from residential_pdf_flow import plan_pages
from notice_word_rows import inspect_rows, rent_table
from docx_patch import patch_docx
from standard_form import inspect_standard
from test_rent_overflow import long_case
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


@pytest.mark.parametrize('count',[7,12,24,60,120,1000])
def test_word_uses_same_page_plan_and_preserves_each_actual_rent_cell(count):
    case=long_case(count);raw=render_residential_notice(case);doc=Document(BytesIO(raw))
    fields=fields_for_case(case);expected=plan_pages(fields,count)
    anchors=[p for p in doc.paragraphs if not p.text and p._p.find('.//'+W+'drawing') is not None]
    assert len(anchors)==len(expected)
    assert sum(bool(p.paragraph_format.page_break_before) for p in anchors)==len(expected)-1
    rows=inspect_rows(doc.element)
    assert len(rows)==count
    assert all(r['period_p'].getparent()!=r['p'].getparent() for r in rows)
    from docx.text.paragraph import Paragraph
    for row in rows:
        assert row['amount']=='$1,000.00'
        for p in [row['p'],row['period_p']]:
            assert all(r.font.name=='Times New Roman' and r.font.size.pt==12 for r in Paragraph(p,doc._body).runs)
    assert saved_word_case(raw)['amounts_due']==[{'due_date':fields[f'PERIOD_{i}'],'amount':'$1,000.00'} for i in range(1,count+1)]


def test_wrapped_and_multiple_page_periods_round_trip_without_added_spaces():
    for label in ['Unpaid rent for March 1 through March 31, 2020 under the rental agreement', 'W'*4000]:
        case=long_case(12);case['amounts_due'][2]['due_date']=label
        raw=render_residential_notice(case)
        assert saved_word_case(raw)['amounts_due'][2]['due_date']==label
        assert len(inspect_rows(Document(BytesIO(raw)).element))==12


def test_rent_edit_contract_accepts_new_tables_and_same_pdf_fields():
    raw=render_residential_notice(long_case(12))
    root=etree.fromstring(zipfile.ZipFile(BytesIO(raw)).read('word/document.xml'))
    table=inspect_standard(root)['rent_tables'][0]
    raw=patch_docx(raw,[{'old':'$1,000.00','new':'$1,250.25','paragraph_index':table['rows'][0]['paragraph_index']},
                        {'old':'$12,000.00','new':'$12,250.25','paragraph_index':table['total']['paragraph_index']}])
    result=render_saved_word_pdf(raw,sha256(raw).hexdigest())
    fields=PdfReader(BytesIO(base64.b64decode(result['pdf_base64']))).get_fields()
    assert fields['AmountDue1']['/V']=='$1,250.25'
    assert fields['TOTAL RENT DUE']['/V']=='$12,250.25'
    assert base64.b64decode(result['pdf_base64'])==render_residential_pdf(saved_word_case(raw))


def test_unbound_table_row_or_rent_moved_after_total_cannot_be_reviewed():
    raw=render_residential_notice(long_case(12))
    for moved in [False,True]:
        doc=Document(BytesIO(raw));table=next(t for t in doc.tables if rent_table(t._tbl))
        if moved:
            table._tbl.getparent().remove(table._tbl)
            doc.element.body.insert(len(doc.element.body)-1,table._tbl)
        else:
            row=table.add_row();row.cells[0].text='Extra period';row.cells[1].text='$10.00'
        out=BytesIO();doc.save(out)
        with pytest.raises(NoticeInputError):saved_word_case(out.getvalue())


def test_editor_and_generator_ship_the_identical_rent_schema_validator():
    assert (ROOT/'notice_word_rows.py').read_bytes()==(ROOT/'revision-service/notice_word_rows.py').read_bytes()


def test_existing_tab_word_is_corrected_after_a_hash_bound_edit(monkeypatch):
    import residential_word_layout
    from residential_word_snapshot import normalize_saved_word
    with monkeypatch.context() as m:
        m.setattr(residential_word_layout,'needs_flow',lambda *_:False)
        legacy=render_residential_notice(long_case(12))
    corrected=normalize_saved_word(legacy,sha256(legacy).hexdigest())
    assert corrected['source_docx_sha256']==sha256(legacy).hexdigest()
    raw=base64.b64decode(corrected['docx_base64'])
    assert saved_word_case(raw)==saved_word_case(legacy)
    assert len(Document(BytesIO(raw)).tables)==2  # rent columns and total
    assert corrected['word_template_version']=='residential-pay-or-quit-2026-10-30-word-layout4'
    with pytest.raises(NoticeInputError):normalize_saved_word(legacy,'0'*64)


def test_longer_period_edit_recalculates_word_page_breaks():
    from residential_word_snapshot import normalize_saved_word
    raw=render_residential_notice(long_case(12))
    root=etree.fromstring(zipfile.ZipFile(BytesIO(raw)).read('word/document.xml'))
    period=inspect_rows(root)[2]
    label='Unpaid rent for March 1 through March 31, 2020 under the rental agreement'
    patched=patch_docx(raw,[{'old':period['period'],'new':label,'paragraph_index':period['period_paragraph_index']}])
    fixed=normalize_saved_word(patched,sha256(patched).hexdigest())
    fixed=base64.b64decode(fixed['docx_base64']);case=saved_word_case(fixed)
    assert case['amounts_due'][2]['due_date']==label
    anchors=[p for p in Document(BytesIO(fixed)).paragraphs if not p.text and p._p.find('.//'+W+'drawing') is not None]
    assert len(anchors)==len(plan_pages(fields_for_case(case),12))==3


def test_normalize_word_endpoint_is_hash_bound_and_read_only(monkeypatch):
    monkeypatch.setenv('DROPBOX_SIGN_API_KEY','local-test-no-network');monkeypatch.setenv('AIRTABLE_API_KEY','local-test-no-network')
    from fastapi.testclient import TestClient
    import main
    raw=render_residential_notice(long_case(12));body={'docx_base64':base64.b64encode(raw).decode(),'expected_sha256':sha256(raw).hexdigest()}
    client=TestClient(main.app);r=client.post('/normalize-saved-notice-word',json=body)
    assert r.status_code==200 and r.json()['is_standard']
    body['expected_sha256']='0'*64
    assert client.post('/normalize-saved-notice-word',json=body).status_code==422


def test_changed_intro_returns_review_hold_not_an_internal_error(monkeypatch):
    monkeypatch.setenv('DROPBOX_SIGN_API_KEY','local-test-no-network');monkeypatch.setenv('AIRTABLE_API_KEY','local-test-no-network')
    from fastapi.testclient import TestClient
    import main
    doc=Document(BytesIO(render_residential_notice(long_case(12))))
    next(p for p in doc.paragraphs if p.text.startswith('PLEASE TAKE NOTICE')).text='Changed introduction'
    out=BytesIO();doc.save(out);raw=out.getvalue()
    r=TestClient(main.app).post('/normalize-saved-notice-word',json={'docx_base64':base64.b64encode(raw).decode(),'expected_sha256':sha256(raw).hexdigest()})
    assert r.status_code==422
