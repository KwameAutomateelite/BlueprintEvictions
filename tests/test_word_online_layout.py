"""Prevent the layout failures observed in Word Online, not only PDF export."""
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from residential_notice import render_residential_notice
from test_residential_standard import case


def test_no_tiny_empty_spacing_paragraphs_that_word_online_expands():
    doc = Document(BytesIO(render_residential_notice(case(6))))
    empty = [p for p in doc.paragraphs if not p.text.strip()]
    # There are only two page artwork anchors; all other gaps use spacing.
    assert len(empty) == 2
    assert all(p.paragraph_format.line_spacing.pt >= 90 for p in empty)
    assert sum(bool(p.paragraph_format.page_break_before) for p in doc.paragraphs) == 1


def test_both_pages_keep_branding_and_footer_with_unique_drawing_ids():
    doc = Document(BytesIO(render_residential_notice(case(3))))
    anchors = [p for p in doc.paragraphs if p._p.find(".//"+qn("w:drawing")) is not None]
    drawing_tag = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'
    drawings = [list(p._p.iter(drawing_tag)) for p in anchors]
    assert len(drawings[0]) >= 2
    assert len(drawings[0]) == len(drawings[1])
    ids = [n.get('id') for nodes in drawings for n in nodes]
    assert len(ids) == len(set(ids))


def test_occupied_and_blank_rent_rows_exceed_the_web_tab_minimum():
    for n in (1, 3, 6):
        doc = Document(BytesIO(render_residential_notice(case(n))))
        rows = [p for p in doc.paragraphs if any(
            b.get(qn('w:name'), '').startswith('BE_RENT_ROW_')
            for b in p._p.iter(qn('w:bookmarkStart')))]
        assert len(rows) == 6
        assert all(p.paragraph_format.line_spacing.pt == 14.5 for p in rows)


def test_signature_rule_is_part_of_printed_signer_not_an_empty_line():
    doc = Document(BytesIO(render_residential_notice(case(1))))
    signer = next(p for p in doc.paragraphs if 'Sam Example' in p.text)
    top = signer._p.find('./'+qn('w:pPr')+'/'+qn('w:pBdr')+'/'+qn('w:top'))
    assert top is not None and top.get(qn('w:val')) == 'single'
    assert top.get(qn('w:sz')) == '14'
    assert 'Date:' not in signer.text
    # Web Word ignores properties after later schema children. Protect order.
    tags = [e.tag for e in signer._p.pPr]
    assert tags.index(qn('w:pBdr')) < tags.index(qn('w:spacing'))
    t = doc.tables[0]
    tags = [e.tag for e in t._tbl.tblPr]
    assert tags.index(qn('w:tblInd')) < tags.index(qn('w:tblLayout'))
    for cell in t.rows[0].cells:
        tags = [e.tag for e in cell._tc.tcPr]
        assert tags.index(qn('w:tcBorders')) < tags.index(qn('w:tcMar'))
