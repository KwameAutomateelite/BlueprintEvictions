"""Meaningful form regressions; no production calls, signing or paid services."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).parents[1]))
import pytest
from docx import Document
from residential_notice import NoticeInputError, fields_for_case, render_residential_notice


def case(rows=3):
    return {
        "notice_type": "residential", "day_count": 3,
        "tenant_names": "Alex Example, Jordan Example, All Other Occupants",
        "property_address_street": "123 Example St, Apt 4",
        "property_address_city": "Sunnyvale, CA 94086", "apt": "4",
        "county": "Santa Clara County",
        "amounts_due": [{"due_date": f"{i + 1}/1/2026", "amount": "1,000"} for i in range(rows)],
        "total_amount_due": str(rows * 1000),
        "landlord_name": "Example Owner LLC",
        "payment_recipient_name": "Example Owner LLC c/o Example Realty",
        "payment_address_street": "456 Example Blvd, Suite 200",
        "payment_address_city": "Cupertino, CA 95014",
        "landlord_phone": "(408) 555-0100", "landlord_email": "owner@example.invalid",
        "signer_name": "Sam Example", "signer_title": "Property Manager",
        "notice_date": "2026-10-26", "service_date": "2026-10-27",
        "case_name": "Example Owner v Example Tenant",
    }


def text_and_parts(data):
    content = render_residential_notice(data)
    doc = Document(BytesIO(content))
    with zipfile.ZipFile(BytesIO(content)) as z:
        parts = {n: z.read(n) for n in z.namelist() if n.startswith("word/")}
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]), parts


def test_semantic_fields_and_static_wording():
    text, _ = text_and_parts(case())
    assert "COUNTY OF SANTA CLARA" in text
    assert "COUNTY OF SANTA CLARA County" not in text
    assert "Apt 4, Apt 4" not in text
    assert "Example Owner LLC c/o Example Realty" in text
    assert "Sam Example, Property Manager" in text
    assert "owner@example.invalid" not in text
    assert "This notice was served to you on _______________." in text
    assert "10/27/26" not in text
    assert "Date: 10/26/26" in text
    assert "TIME PERIOD DUE\tAMOUNT DUE" in text
    assert "$3,000.00" in text
    assert "{{" not in text


def test_template_structure_is_identical_for_one_three_and_six_rows():
    # Filling different cases must not move/remove paragraphs or rewrite any
    # header, footer, artwork, section, style or relationship package part.
    import lxml.etree as ET
    all_parts = [text_and_parts(case(n))[1] for n in (1, 3, 6)]
    for name, original in all_parts[0].items():
        if name != "word/document.xml":
            assert all(p[name] == original for p in all_parts[1:]), name
    layouts = []
    for parts in all_parts:
        root = ET.fromstring(parts["word/document.xml"])
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for t in root.findall(".//w:t", ns):
            t.getparent().remove(t)
        layouts.append(ET.tostring(root))
    assert layouts[0] == layouts[1] == layouts[2]


@pytest.mark.parametrize("change", [
    {"total_amount_due": "2999.99"},
    {"amounts_due": [{"due_date": "8/1/2026", "amount": "NOT_FOUND"}]},
    {"tenant_names": "WideName " * 35},
    {"county": ""},
    {"notice_date": "2026-02-30"},
    {"notice_header": "THREE DAY BREACH NOTICE"},
    {"payment_recipient_name": "Owners name: Example Owner, Company: Example LLC"},
])
def test_incomplete_inconsistent_or_overflowing_input_is_held(change):
    data = case(); data.update(change)
    with pytest.raises(NoticeInputError):
        render_residential_notice(data)


def test_unused_rows_stay_blank_and_stray_separator_is_removed():
    data = case(1); data["amounts_due"][0]["due_date"] = "; ; 8/15/2026"
    fields = fields_for_case(data)
    assert fields["PERIOD_1"] == "8/15/2026"
    assert fields["PERIOD_2"] == fields["AMOUNT_6"] == ""


def test_missing_signer_is_not_replaced_by_owner_or_stale_reference_person():
    data = case(); data["signer_name"] = data["signer_title"] = ""
    with pytest.raises(NoticeInputError, match="Designated intake signer's"):
        render_residential_notice(data)
    data["signer_name"] = "Sam Example"
    with pytest.raises(NoticeInputError):
        render_residential_notice(data)



def test_generate_endpoint_preserves_docx_and_attachment_headers():
    import os
    os.environ.setdefault("DROPBOX_SIGN_API_KEY", "dummy")
    os.environ.setdefault("AIRTABLE_API_KEY", "dummy")
    from fastapi.testclient import TestClient
    from main import app
    data = case(); data["attachments_required"] = ["Sample required attachment"]
    with TestClient(app) as client:
        response = client.post("/generate-notice", json=data)
        assert response.status_code == 200
        assert response.content[:2] == b"PK"
        assert "wordprocessingml.document" in response.headers["content-type"]
        assert response.headers["x-attachments-needed"] == "Sample required attachment"
        assert response.headers["x-notice-template-version"] == "residential-pay-or-quit-2026-10-26-layout2"
        data["total_amount_due"] = "1"
        held = client.post("/generate-notice", json=data)
        assert held.status_code == 422
        assert "sum" in held.json()["detail"]
