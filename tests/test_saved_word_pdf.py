from hashlib import sha256
from io import BytesIO
from pathlib import Path
import base64
import sys

import pytest
from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).parents[1]))
sys.path.insert(0, str(Path(__file__).parents[1] / "revision-service"))
sys.path.insert(0, str(Path(__file__).parent))
from residential_notice import NoticeInputError, render_residential_notice
from residential_pdf import render_residential_pdf
from residential_word_snapshot import render_saved_word_pdf, saved_word_case, _bindings
from test_residential_standard import case
from docx_patch import patch_docx
from standard_form import inspect_standard
from lxml import etree
import zipfile


def snapshot(raw):
    return render_saved_word_pdf(raw, sha256(raw).hexdigest())


@pytest.mark.parametrize("rows", [1, 3, 6])
def test_saved_word_produces_same_pdf_as_original_case(rows):
    data = case(rows)
    raw = render_residential_notice(data)
    result = snapshot(raw)
    assert result["is_standard"] is True
    pdf = base64.b64decode(result["pdf_base64"])
    assert pdf == render_residential_pdf(data)
    assert result["pdf_sha256"] == sha256(pdf).hexdigest()
    assert snapshot(raw) == result


def test_saved_editor_rent_revision_changes_both_pdf_columns_and_total():
    raw = render_residential_notice(case())
    root = etree.fromstring(zipfile.ZipFile(BytesIO(raw)).read("word/document.xml"))
    table = inspect_standard(root)["rent_tables"][0]
    changes = [{"old": table["rows"][0]["text"], "new": "$1,250.25", "paragraph_index": table["rows"][0]["paragraph_index"]},
               {"old": table["total"]["text"], "new": "$3,250.25", "paragraph_index": table["total"]["paragraph_index"]}]
    edited = patch_docx(raw, changes)
    fields = PdfReader(BytesIO(base64.b64decode(snapshot(edited)["pdf_base64"]))).get_fields()
    assert fields["Month1"]["/V"] == "January 1-31, 2026"
    assert fields["AmountDue1"]["/V"] == "$1,250.25"
    assert fields["TOTAL RENT DUE"]["/V"] == "$3,250.25"
    assert snapshot(edited)["pdf_sha256"] != snapshot(raw)["pdf_sha256"]


def test_saved_county_and_signer_edits_are_not_overwritten_by_cached_inputs():
    raw = render_residential_notice(case())
    edited = patch_docx(raw, [{"old": "COUNTY OF SANTA CLARA", "new": "COUNTY OF SAN MATEO"},
                             {"old": "Sam Example", "new": "Pat Example"}])
    fields = PdfReader(BytesIO(base64.b64decode(snapshot(edited)["pdf_base64"]))).get_fields()
    assert fields["COUNTY OF"]["/V"] == "SAN MATEO"
    assert fields["Text5"]["/V"] == "Pat Example, Property Manager"
    assert fields["PLAINTIFF"]["/V"] == "Example Owner LLC c/o Example Realty"


def test_changed_source_and_modified_boilerplate_cannot_be_approved():
    raw = render_residential_notice(case())
    with pytest.raises(NoticeInputError):
        render_saved_word_pdf(raw, "0" * 64)
    doc = Document(BytesIO(raw))
    for p in doc.paragraphs:
        if "Please be advised" in p.text:
            p.text = "Changed legal content"
            break
    out = BytesIO(); doc.save(out)
    with pytest.raises(NoticeInputError, match="boilerplate"):
        snapshot(out.getvalue())


def test_other_notice_families_keep_existing_pdf_conversion():
    doc = Document(); doc.add_paragraph("Different notice family")
    out = BytesIO(); doc.save(out)
    assert snapshot(out.getvalue())["is_standard"] is False


def test_snapshot_api_hash_and_invalid_inputs(monkeypatch):
    monkeypatch.setenv("DROPBOX_SIGN_API_KEY", "local-test-no-network")
    monkeypatch.setenv("AIRTABLE_API_KEY", "local-test-no-network")
    from fastapi.testclient import TestClient
    import main
    raw = render_residential_notice(case())
    payload = {"docx_base64": base64.b64encode(raw).decode(), "expected_sha256": sha256(raw).hexdigest()}
    client = TestClient(main.app)
    result = client.post("/render-saved-notice-pdf", json=payload)
    assert result.status_code == 200 and result.json()["is_standard"]
    payload["expected_sha256"] = "0" * 64
    assert client.post("/render-saved-notice-pdf", json=payload).status_code == 422
    payload["docx_base64"] = "not valid base64"
    assert client.post("/render-saved-notice-pdf", json=payload).status_code == 422
