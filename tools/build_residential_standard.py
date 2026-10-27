"""Build the blank, editable Word master from the approved October 2026 form.

The retained legacy package supplies only its existing Blueprint artwork and
footer. All case slots are blank placeholders, never values from the example.
Run locally when the approved master changes; generation does not run this tool.
"""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "templates/3-Day Notice - BLUEPRINT - residential - NEW BRANDING_071125_v2.docx"
TARGET = ROOT / "templates/residential-pay-or-quit-2026-10.docx"


def build():
    doc = Document(SOURCE)
    artwork = deepcopy(doc.paragraphs[0]._p)
    footer_part = doc.sections[0].footer.part
    footer_run = deepcopy(footer_part.element.find(".//" + qn("w:r")))
    for blip in footer_run.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        rid = blip.get(qn("r:embed"))
        blip.set(qn("r:embed"), doc.part.relate_to(footer_part.rels[rid].target_part, RT.IMAGE))
    sect = deepcopy(doc.sections[-1]._sectPr)
    for node in list(doc._element.body):
        doc._element.body.remove(node)
    doc._element.body.append(sect)
    for ref in list(sect.findall(qn("w:footerReference"))):
        sect.remove(ref)
    pg = sect.find(qn("w:pgMar"))
    pg.set(qn("w:bottom"), "720")
    # Page geometry/artwork/footer are inherited from the source Word package.
    doc._element.body.insert(0, artwork)
    # The reference footer is artwork anchored to its first page. Put that
    # artwork on the retained page anchor so Word/LibreOffice agree even when
    # a footer's tiny inherited paragraph height is collapsed.
    artwork.append(deepcopy(footer_run))
    from docx.text.paragraph import Paragraph
    ap = Paragraph(artwork, doc._body)
    ap.paragraph_format.space_before = Pt(0)
    ap.paragraph_format.space_after = Pt(0)
    # Word Online clamps tiny empty paragraphs to about 7.1 pt even with an
    # exact line height. One full-height anchor replaces all leading spacers.
    ap.paragraph_format.line_spacing = Pt(98.2)
    cursor = 98.2

    def raise_runs(paragraph, points):
        # PDF form appearances can place text above a normal Word baseline.
        # Character position preserves paragraph geometry while matching ink.
        for run in paragraph.runs:
            position = OxmlElement("w:position")
            position.set(qn("w:val"), str(round(points*2)))
            run._r.get_or_add_rPr().insert_element_before(position, "w:sz", "w:szCs", "w:u", "w:lang")

    def block(text, y, *, x=72, width=480, size=11.5, leading=13.75,
              bold=False, bold_prefix=None, underline=False, align=None,
              slot=None, new_page=False, raise_by=0):
        nonlocal cursor
        if new_page:
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.line_spacing = Pt(y)
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            # The approved form repeats the same brand band and footer on
            # both pages; drawing anchors do not consume paragraph height.
            for child in artwork:
                if child.tag != qn("w:pPr"):
                    p._p.append(deepcopy(child))
            cursor = y
        gap = y - cursor
        if gap < -0.1:
            raise ValueError(f"Overlapping master block: {slot or text[:30]} ({gap})")
        p = doc.add_paragraph(style="Normal")
        pf = p.paragraph_format
        pf.left_indent = Pt(x - 54)
        pf.right_indent = Pt(594 - x - width)
        pf.first_line_indent = Pt(0)
        pf.space_before, pf.space_after = Pt(max(0, gap)), Pt(0)
        pf.line_spacing = Pt(leading)
        pf.keep_together = True
        pf.keep_with_next = False
        pf.widow_control = False
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
        if slot:
            bookmark = OxmlElement("w:bookmarkStart")
            bookmark.set(qn("w:id"), str(len(doc.paragraphs)))
            bookmark.set(qn("w:name"), "BE_" + slot)
            p._p.append(bookmark)
        chunks = [(text, bold)]
        if bold_prefix and text.startswith(bold_prefix):
            chunks = [(bold_prefix, True), (text[len(bold_prefix):], False)]
        for value, weight in chunks:
            r = p.add_run(value)
            r.bold = weight
            r.underline = underline
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0, 0, 0)
        if slot:
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), bookmark.get(qn("w:id")))
            p._p.append(end)
        if raise_by:
            raise_runs(p, raise_by)
        cursor = y + leading * (text.count("\n") + 1)
        return p

    block("THREE (3) DAY NOTICE TO PAY OR QUIT", 98.2, size=15.5, leading=20,
          bold=True, x=68.3, width=480, align=WD_ALIGN_PARAGRAPH.CENTER)
    tenant = block("To: {{TENANT_NAMES}}\n{{TENANT_NAMES_CONTINUED}}", 130, size=12, leading=14, slot="TENANT_NAMES")
    tenant.runs[0].text = "To: "
    tenant.runs[0].font.size = Pt(11.5)
    field = tenant.add_run("{{TENANT_NAMES}}\n{{TENANT_NAMES_CONTINUED}}")
    field.font.name, field.font.size = "Times New Roman", Pt(12)
    block("Tenants in possession of the premises herein referred to are situated at:",
          158, leading=14, raise_by=1)
    block("{{PROPERTY_STREET}}", 180, x=106, width=445, size=12, leading=14, slot="PROPERTY_STREET", raise_by=1)
    block("{{PROPERTY_CITY}}", 194, x=106, width=445, size=12, leading=14, slot="PROPERTY_CITY")
    county = block("COUNTY OF {{COUNTY}}", 212, x=108, width=445, size=12, leading=14, slot="COUNTY")
    county.runs[0].text = "COUNTY OF "
    county.runs[0].font.size = Pt(11.5)
    field = county.add_run("{{COUNTY}}")
    field.font.name, field.font.size = "Times New Roman", Pt(12)
    raise_runs(county, -0.5)
    block('[hereinafter called the "Premises"] and all notices thereafter:', 227, leading=14, raise_by=0.5)
    please = block("PLEASE TAKE NOTICE that pursuant to the terms of tenancy for the above-referenced\n"
          "agreement by which you hold the possession of the Premises described in this Notice, there is\n"
          "now due, unpaid, and delinquent rent in the following amounts as specified below:",
          254, leading=15.8, bold_prefix="PLEASE TAKE NOTICE")
    # The approved PDF's first three lines use different vertical offsets.
    # Keep their paragraph geometry fixed and align each printed baseline.
    for run in list(please.runs):
        please._p.remove(run._r)
    for text, bold, offset in [
        ("PLEASE TAKE NOTICE", True, 1.5),
        (" that pursuant to the terms of tenancy for the above-referenced\n", False, 1.5),
        ("agreement by which you hold the possession of the Premises described in this Notice, there is\n", False, 3),
        ("now due, unpaid, and delinquent rent in the following amounts as specified below:", False, 1)]:
        run = please.add_run(text)
        run.font.name, run.font.size, run.bold = "Times New Roman", Pt(11.5), bold
        position=OxmlElement("w:position");position.set(qn("w:val"),str(round(offset*2)))
        run._r.get_or_add_rPr().insert_element_before(position, "w:sz", "w:u")
    # 14.5pt exceeds Word Online's fallback tab glyph height (~14.37pt),
    # keeping occupied and blank rows at the same fixed height.
    p = block("TIME PERIOD DUE\tAMOUNT DUE", 311, x=180, width=280,
              size=11.5, leading=14.5, bold=True, underline=True)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(306))
    p.runs[0].text = "TIME PERIOD DUE"
    tab = p.add_run("\t")
    tab.font.name, tab.font.size, tab.underline = "Times New Roman", Pt(11.5), False
    amount_header = p.add_run("AMOUNT DUE")
    amount_header.font.name, amount_header.font.size = "Times New Roman", Pt(11.5)
    amount_header.bold = amount_header.underline = True
    raise_runs(p, 0.5)
    for i in range(1, 7):
        p = block("{{PERIOD_%s}}\t{{AMOUNT_%s}}" % (i, i), 327 + (i - 1) * 14.5,
                  x=180, width=280, size=12, leading=14.5, slot="RENT_ROW_%s" % i)
        p.paragraph_format.tab_stops.add_tab_stop(Pt(306))
    p = block("", 414, x=180, width=292, size=12, leading=16)
    Paragraph(p._p.getprevious(), doc._body).paragraph_format.space_after = p.paragraph_format.space_before
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Pt(168), Pt(118)
    props = table._tbl.tblPr
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "2540")
    indent.set(qn("w:type"), "dxa")
    props.insert(1, indent)
    table.rows[0].height = Pt(16)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Pt(168 if i == 0 else 118)
        cp = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for edge in ("top", "bottom", "left", "right"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:w"), ("140" if i else "40") if edge in ("left", "right") else "0")
            e.set(qn("w:type"), "dxa")
            margins.append(e)
        cp.append(margins)
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "bottom", "left", "right"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single" if i else "nil")
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), "000000")
            borders.append(e)
        cp.insert(1, borders)
        para = cell.paragraphs[0]
        pf = para.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Pt(14)
        pf.keep_with_next = False
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run("TOTAL AMOUNT DUE:" if i == 0 else "{{TOTAL_AMOUNT}}")
        run.font.name, run.font.size, run.bold = "Times New Roman", Pt(11.5 if i == 0 else 12), True
        raise_runs(para, 1)
    p._p.addprevious(table._tbl)
    p._p.getparent().remove(p._p)
    block("You are further notified that within three (3) days after service of this notice on you, not including\n"
          "Saturdays, Sundays or other judicial holidays, you must pay the total amount of rent stated in this\n"
          "Notice in full or quit the Premises, meaning to deliver up possession of the Premises and vacate.",
          439, x=73.7, width=482, leading=13.75, raise_by=1)
    service = block("This notice was served to you on _______________. You have three (3) days after the date the\n"
          "notice was served to you, excluding Saturdays, Sundays, and other judicial holidays, to pay the\n"
          "rent due or vacate the property.", 492, x=73.7, width=482, leading=13.75, bold=True)
    before, after = service.runs[0].text.split("days after", 1)
    service.runs[0].text = before
    for value, under in [("days after", True), (after, False)]:
        r = service.add_run(value)
        r.font.name, r.font.size, r.bold, r.underline = "Times New Roman", Pt(11.5), True, under
    raise_runs(service, 0.5)
    block("Payment of the above total amount due should be tendered to, or possession of the premises\n"
          "delivered to:", 543.5, x=73.7, width=482, leading=13.75, raise_by=1)
    block("{{PAYMENT_RECIPIENT}}", 571, x=157, width=420, size=12, leading=14, slot="PAYMENT_RECIPIENT", raise_by=3.5)
    block("{{PAYMENT_STREET}}", 586, x=159, width=420, size=12, leading=14, slot="PAYMENT_STREET", raise_by=2)
    block("{{PAYMENT_CITY}}", 600, x=159, width=420, size=12, leading=14, slot="PAYMENT_CITY", raise_by=1)
    block("{{PAYMENT_PHONE}}", 616, x=160, width=415, size=12, leading=14, slot="PAYMENT_PHONE", raise_by=1.5)
    block("Rent or possession of the Premises may be tendered between normal business hours of 9:00AM to\n"
          "5:00PM, Monday through Friday.", 630, x=73.7, width=482, leading=13.75, raise_by=2)
    payment_terms = block("Only the total amount due and made prior to the expiration of this Notice may be made using cash or\n"
          "certified funds (e.g., cashier's check, money order, etc.). No other forms of payment are acceptable.\n"
          "Landlords, and the attorney on behalf of the Landlords, will not accept any payments after the\n"
          "expiration of this Notice. Further, be advised that Landlords will not accept any partial payments for the\n"
          "amounts outlined in this notice nor will they accept payments made to others ostensibly on their behalf:\n"
          "all payments must be made to the above-stated address.", 670.5, x=73.7, width=482, leading=13.75)
    final_line = "all payments must be made to the above-stated address."
    payment_terms.runs[0].text = payment_terms.runs[0].text[:-len(final_line)]
    final_run = payment_terms.add_run(final_line)
    final_run.font.name = "Times New Roman"
    final_run.font.size = Pt(11.5)
    final_run.underline = True
    raise_runs(payment_terms, 1.5)
    block("You are further notified by this Notice the Landlords elect to and therefore do declare a forfeiture of\n"
          "the rental agreement if the rent stated in this notice is not paid in full within the 3 (three) days.",
          98, leading=13.75, new_page=True)
    block("NOTICES FROM THE STATE OF CALIFORNIA:", 140.5, x=78.4, width=475, bold=True)
    block("(A) NOTICE regarding possible credit consequences: Pursuant to Civil Code § 1785.26, you are\n"
          "hereby notified that a negative credit report reflecting on your credit record may be submitted in the\n"
          "future to a credit reporting agency if you fail to fulfill the terms of your rental obligation or if you\n"
          "default in those obligations in any way. This is the only notice that you will receive in this regard.",
          167, leading=13.75, bold_prefix="(A) NOTICE regarding possible credit consequences:")
    block('(B) NOTICE regarding abandoned property: "State law permits former tenants to reclaim\n'
          'abandoned personal property left at the former address of the tenant, subject to certain conditions.\n'
          'You may or may not be able to reclaim property without incurring additional costs, depending on\n'
          'the cost of storing the property and the length of time before it is reclaimed. In general, these costs\n'
          'will be lower the sooner you contact your former landlord after being notified that property\n'
          'belonging to you was left behind after you moved out." California Civil Codes §1946.1; §1983-\n'
          '1989 et seq.', 236, leading=13.75, bold_prefix="(B) NOTICE regarding abandoned property:")
    block("Please be advised that any tenant under the above-stated agreement that remains in the unit after the\n"
          "termination date may result in the Landlords seeking to enforce the termination in Court, in an action\n"
          "for Unlawful Detainer at which time the occupant(s) may appear and present a defense.",
          346, leading=13.75)
    block("Finally, any Notice to Pay or Quit served prior to this one is hereby superseded and canceled. Any notice\n"
          "served in conjunction with or subsequent to this one is not designed to supersede or cancel this present\n"
          "notice. You are still required to oblige by the terms outlined in this Notice.",
          401, width=486, leading=13.75)
    # Date/caption is a real text paragraph; the rule belongs to the separate
    # signer paragraph. No hanging line, table or empty spacer is needed.
    p = block("Date: {{NOTICE_DATE}}\t(Original signed by agent/owner)",
              484, x=77.65, width=385.82, size=11.5, leading=14,
              slot="NOTICE_DATE")
    p.paragraph_format.tab_stops.add_tab_stop(Pt(194.8))
    p.runs[0].text = "Date: "
    value = p.add_run("{{NOTICE_DATE}}")
    value.font.name, value.font.size = "Times New Roman", Pt(12)
    caption = p.add_run("\t(Original signed by agent/owner)")
    caption.font.name, caption.font.size = "Times New Roman", Pt(11.5)
    signer = block("{{SIGNER_NAME_AND_TITLE}}\n{{SIGNER_CONTINUED}}", 498, x=193.4, width=268.65, size=12,
          leading=14, align=WD_ALIGN_PARAGRAPH.CENTER, slot="SIGNER_NAME_AND_TITLE")
    # The top border belongs to the signer's paragraph, so it spans only
    # the signature area and follows the same text flow in Word and PDF.
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement("w:top")
    for key, val in {"val":"single", "sz":"14", "space":"5", "color":"000000"}.items():
        edge.set(qn("w:"+key), val)
    borders.append(edge)
    signer._p.get_or_add_pPr().insert_element_before(borders, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc")
    cursor += 6.75  # 1.75pt rule plus 5pt clearance, outside the 28pt text box.
    block("Housing Provider's Address:", 540, x=216.4, width=365, size=11.5, leading=14)
    block("{{PROVIDER_STREET}}", 556, x=219, width=375, size=12, leading=14, slot="PROVIDER_STREET")
    block("{{PROVIDER_CITY}}", 571, x=217, width=375, size=12, leading=14, slot="PROVIDER_CITY")
    phone = block("Phone: {{PROVIDER_PHONE}}", 585.5, x=216.65, width=365, size=12, leading=14, slot="PROVIDER_PHONE")
    phone.runs[0].text = "Phone: "
    phone.runs[0].font.size = Pt(11.5)
    value = phone.add_run("{{PROVIDER_PHONE}}")
    value.font.name, value.font.size = "Times New Roman", Pt(12)
    doc.core_properties.title = "Blueprint residential three day notice to pay or quit"
    doc.core_properties.subject = "Approved layout 2026-10-26; case fields only"
    doc.core_properties.comments = "Reference PDF SHA256 0e23f57004a2f9842acaab92a334d46c031a5d09e31f4f3900b91f531418825d"
    # Drawing IDs are unique across the two retained page artworks.
    for i, prop in enumerate(doc._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'), 1):
        prop.set("id", str(i))
    doc.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(build())
